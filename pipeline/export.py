"""Excel export — the tracker described in SOP Module 63.

Two files are written per run:
  companies_<label>_<stamp>_<runid>.xlsx      full cumulative workbook (below)
  new_companies_<label>_<stamp>_<runid>.xlsx  ONLY companies first seen this
                                               run — the daily new-companies
                                               deliverable, no repeats/history

Workbook sheets (companies_* file):
  Companies        one row per approved/review account, tracker field order
  New_This_Run     the same rows, filtered to companies first seen in this
                    run — so a daily run surfaces what's actually new without
                    re-scanning the full cumulative tracker for repeats
  Decision_Makers  empty scaffold with the Module 44 columns, ready for phase 2
  Evidence         source trail: company, source type, URL, fact, date checked
  Signals          every detected signal with the page it came from
  Excluded         the exclusion ledger, so nobody re-researches these
  Run_Report       quality metrics and distributions for this run
  Learning_Log     observations captured during the run (Module 15)
"""

from __future__ import annotations

import json
import logging
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .store import Store

log = logging.getLogger("pipeline.export")

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TIER_FILLS = {
    "A": PatternFill("solid", fgColor="C6EFCE"),
    "B": PatternFill("solid", fgColor="FFEB9C"),
    "C": PatternFill("solid", fgColor="F2F2F2"),
}

# (db column, header, width) in SOP Module 63 field-group order.
COMPANY_COLUMNS: list[tuple[str, str, int]] = [
    ("id", "ID", 6),
    ("tier", "Tier", 7),
    ("score_total", "Score", 8),
    ("status", "Status", 12),
    # Company identity
    ("company_name", "Company Name", 34),
    ("website", "Website", 34),
    ("linkedin_url", "LinkedIn URL", 30),
    ("headquarters_city", "HQ City", 18),
    ("headquarters_state", "HQ State", 9),
    ("hq_confirmed", "HQ Confirmed", 13),
    ("discovery_region", "Searched Region", 15),
    ("country", "Country", 14),
    ("phone", "Phone", 16),
    ("employee_range", "Employee Range", 15),
    ("employee_estimate", "Employee Estimate", 15),
    ("employee_range_source", "Employee Range Source", 18),
    # Company model
    ("entity_type", "Entity Type", 18),
    ("entity_confidence", "Entity Confidence", 16),
    ("llm_entity_type", "LLM Entity Type", 17),
    ("llm_agrees", "LLM Agrees", 11),
    ("llm_challenge_verdict", "LLM Challenge", 14),
    ("llm_challenge_notes", "LLM Concerns", 45),
    ("is_us", "Is US", 8),
    ("segment", "Segment", 28),
    ("services", "Services", 45),
    ("recurring_evidence", "Recurring-Service Evidence", 40),
    ("customer_industries", "Customer Industries", 30),
    ("company_description", "Company Description", 55),
    # Qualification
    ("fit_reason", "Fit Reason", 70),
    ("operational_signals", "Operational Signals", 45),
    ("intent_score", "Intent Score (20)", 15),
    ("intent_summary", "Intent Summary", 55),
    ("open_roles", "Open Ops Roles", 14),
    ("ats_provider", "ATS", 18),
    ("tool_categories", "Tool Categories", 14),
    ("tools_found", "Tools Found", 45),
    ("timing_signal", "Timing Signal", 45),
    ("exclusions_checked", "Exclusions Checked", 30),
    ("score_fit", "Score: Fit (40)", 13),
    ("score_need", "Score: Need (25)", 14),
    ("score_timing", "Score: Timing (20)", 15),
    ("score_access", "Score: Access (15)", 15),
    ("status_reason", "Status Reason", 45),
    # Workflow
    ("owner", "Owner", 12),
    ("source_pipeline", "Source Pipeline", 16),
    ("discovery_source", "Discovery Source", 30),
    ("discovery_query", "Discovery Query", 40),
    ("name_source", "Name Source", 16),
    ("rationale_source", "Rationale Source", 16),
    ("geo_reason", "Geo Evidence", 30),
    ("next_action", "Next Action", 40),
    ("review_notes", "Review Notes", 30),
    ("duplicate_status", "Duplicate Status", 18),
    ("clean_verified", "Clean Verified", 14),
    ("cleaned_at", "Cleaned At", 20),
    ("completeness_pct", "Completeness %", 14),
    ("times_seen", "Times Seen", 11),
    ("first_seen_at", "First Seen", 20),
    ("last_updated_at", "Last Updated", 20),
]

DECISION_MAKER_COLUMNS: list[tuple[str, int]] = [
    ("Company ID", 10), ("Company Name", 34), ("Person Name", 24), ("Title", 30),
    ("Role Type", 18), ("Profile URL", 34), ("Location", 20), ("Relevance", 45),
    ("Recent Activity", 30), ("Verified Email", 26), ("Source URL", 34),
    ("Date Checked", 14), ("Notes", 30),
]


def export_workbook(store: Store, cfg: Any, run_id: str, report: dict[str, Any]) -> Path:
    out_dir = Path(cfg.get("export.output_dir", "data/exports"))
    if not out_dir.is_absolute():
        out_dir = Path(__file__).resolve().parent.parent / out_dir
    # One folder per run date, so a day's output stays together.
    if cfg.get("export.date_folders", True):
        out_dir = out_dir / datetime.now().strftime("%Y-%m-%d")
    out_dir.mkdir(parents=True, exist_ok=True)

    # Unique, dated filename: label + date + time + run id, so two runs on the
    # same day never collide and every file traces back to its run log.
    label = cfg.get("run.label", "run")
    now = datetime.now()
    stamp = now.strftime("%Y-%m-%d_%H%M%S")
    short_run = run_id.replace("run_", "")[-6:]
    path = out_dir / f"companies_{label}_{stamp}_{short_run}.xlsx"
    counter = 1
    while path.exists():
        path = out_dir / f"companies_{label}_{stamp}_{short_run}_{counter}.xlsx"
        counter += 1

    wb = Workbook()
    wb.remove(wb.active)

    statuses = cfg.get("export.include_statuses", ["approved", "review"])
    tiers = cfg.get("export.include_tiers", ["A", "B", "C"])

    placeholders = ", ".join("?" for _ in statuses)
    tier_marks = ", ".join("?" for _ in tiers)

    # Only post-clean records reach the workbook. `clean_verified` is stamped
    # solely by the clean stage, so a demoted or unverified row cannot appear
    # even if its status still looks acceptable.
    verified_only = bool(cfg.get("export.verified_only", True))
    clean_clause = " AND clean_verified = 1" if verified_only else ""

    rows = store.companies_where(
        f"status IN ({placeholders}) AND (tier IN ({tier_marks}) OR tier IS NULL)"
        f"{clean_clause}"
        " ORDER BY CASE tier WHEN 'A' THEN 1 WHEN 'B' THEN 2 WHEN 'C' THEN 3 ELSE 4 END,"
        " score_total DESC",
        (*statuses, *tiers),
    )
    if verified_only and not rows:
        log.warning(
            "No clean-verified records to export. Did the clean stage run? "
            "Set export.verified_only=false to export unverified rows."
        )

    new_rows = [r for r in rows if r["first_seen_run"] == run_id]

    _sheet_companies(wb, rows, cfg)
    if bool(cfg.get("export.new_this_run_sheet", True)):
        _sheet_companies(wb, new_rows, cfg, sheet_name="New_This_Run")
        log.info(
            "  %d of %d exported companies are new this run (first_seen_run=%s)",
            len(new_rows), len(rows), run_id,
        )
    _sheet_decision_makers(wb, rows)
    _sheet_evidence(wb, store, rows)
    _sheet_signals(wb, store, rows)
    _sheet_excluded(wb, store)
    _sheet_run_report(wb, report, store)
    _sheet_learning(wb, store, run_id)

    wb.save(path)
    log.info("Exported %d companies to %s", len(rows), path)

    if bool(cfg.get("export.new_companies_workbook", True)):
        new_path = _export_new_companies_only(new_rows, cfg, out_dir, label, stamp, short_run)
        log.info("Exported %d new companies to %s", len(new_rows), new_path)

    return path


def _export_new_companies_only(
    new_rows: list[Any], cfg: Any, out_dir: Path, label: str, stamp: str, short_run: str,
) -> Path:
    """A standalone daily workbook containing ONLY companies first seen in this
    run — no cumulative history, so there is never a stale/repeat row to
    mistake for something new. The main workbook's New_This_Run tab holds the
    same rows; this file exists for handoff to people who should only ever
    see today's net-new accounts."""
    path = out_dir / f"new_companies_{label}_{stamp}_{short_run}.xlsx"
    counter = 1
    while path.exists():
        path = out_dir / f"new_companies_{label}_{stamp}_{short_run}_{counter}.xlsx"
        counter += 1

    wb = Workbook()
    wb.remove(wb.active)
    _sheet_companies(wb, new_rows, cfg, sheet_name="New Companies")
    wb.save(path)
    return path


# --------------------------------------------------------------------------


def _style_header(ws: Worksheet, widths: list[int], freeze: bool = True, autofilter: bool = True) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    if freeze:
        ws.freeze_panes = "A2"
    if autofilter and ws.max_row >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(widths))}{max(ws.max_row, 1)}"


def _sheet_companies(wb: Workbook, rows: list[Any], cfg: Any, sheet_name: str = "Companies") -> None:
    ws = wb.create_sheet(sheet_name)
    ws.append([header for _, header, _ in COMPANY_COLUMNS])

    for row in rows:
        ws.append([_cell(row, key) for key, _, _ in COMPANY_COLUMNS])

    # Colour-code the tier column.
    tier_col = next(i for i, (k, _, _) in enumerate(COMPANY_COLUMNS, start=1) if k == "tier")
    for excel_row in range(2, ws.max_row + 1):
        cell = ws.cell(row=excel_row, column=tier_col)
        fill = TIER_FILLS.get(str(cell.value))
        if fill:
            cell.fill = fill
        cell.alignment = Alignment(horizontal="center")

    # Hyperlink the website column.
    web_col = next(i for i, (k, _, _) in enumerate(COMPANY_COLUMNS, start=1) if k == "website")
    for excel_row in range(2, ws.max_row + 1):
        cell = ws.cell(row=excel_row, column=web_col)
        if cell.value and str(cell.value).startswith("http"):
            cell.hyperlink = str(cell.value)
            cell.font = Font(color="0563C1", underline="single")

    _style_header(
        ws,
        [width for _, _, width in COMPANY_COLUMNS],
        freeze=bool(cfg.get("export.freeze_header", True)),
        autofilter=bool(cfg.get("export.autofilter", True)),
    )


def _sheet_decision_makers(wb: Workbook, rows: list[Any]) -> None:
    """Scaffold for phase 2 — pre-seeded with the approved Tier A/B accounts."""
    ws = wb.create_sheet("Decision_Makers")
    ws.append([header for header, _ in DECISION_MAKER_COLUMNS])

    for row in rows:
        if row["status"] != "approved" or row["tier"] not in ("A", "B"):
            continue
        target = 4 if row["tier"] == "A" else 3  # SOP Module 31 depth targets
        for _ in range(target):
            ws.append([row["id"], row["company_name"], "", "", "", "", "", "", "", "", "", "", ""])

    _style_header(ws, [width for _, width in DECISION_MAKER_COLUMNS])


def _sheet_evidence(wb: Workbook, store: Store, rows: list[Any]) -> None:
    ws = wb.create_sheet("Evidence")
    ws.append([
        "Company ID", "Company Name", "Evidence Level", "Source Type",
        "Source URL", "Supported Fact", "Date Checked",
    ])
    by_id = {r["id"]: r["company_name"] for r in rows}
    if by_id:
        marks = ", ".join("?" for _ in by_id)
        records = store.conn.execute(
            f"SELECT * FROM evidence WHERE company_id IN ({marks})"
            " ORDER BY company_id, evidence_level",
            tuple(by_id),
        ).fetchall()
        for rec in records:
            ws.append([
                rec["company_id"], by_id.get(rec["company_id"], ""), rec["evidence_level"],
                rec["source_type"], rec["source_url"], rec["supported_fact"], rec["date_checked"],
            ])
    _style_header(ws, [10, 30, 14, 22, 50, 60, 14])


def _sheet_signals(wb: Workbook, store: Store, rows: list[Any]) -> None:
    ws = wb.create_sheet("Signals")
    ws.append(["Company ID", "Company Name", "Category", "Signal", "Detail", "Source URL"])
    by_id = {r["id"]: r["company_name"] for r in rows}
    if by_id:
        marks = ", ".join("?" for _ in by_id)
        records = store.conn.execute(
            f"SELECT * FROM signals WHERE company_id IN ({marks}) ORDER BY company_id, category",
            tuple(by_id),
        ).fetchall()
        for rec in records:
            ws.append([
                rec["company_id"], by_id.get(rec["company_id"], ""), rec["category"],
                rec["signal"], rec["detail"], rec["source_url"],
            ])
        # Dated timing events (Change 1) live in intent_events, not signals —
        # fold them in here so this sheet stays the complete signal picture.
        # The date suffix shows at a glance whether a timing entry carries a
        # real event_date or fell back to the undated confidence penalty.
        events = store.conn.execute(
            f"SELECT * FROM intent_events WHERE company_id IN ({marks}) ORDER BY company_id",
            tuple(by_id),
        ).fetchall()
        for rec in events:
            when = f"dated {rec['event_date'][:10]}" if rec["event_date"] else "undated"
            detail = f"{rec['detail'] or ''} [{when}]".strip()
            ws.append([
                rec["company_id"], by_id.get(rec["company_id"], ""), "timing",
                rec["event_type"], detail, rec["source_url"],
            ])
    _style_header(ws, [10, 30, 16, 24, 60, 45])


def _sheet_excluded(wb: Workbook, store: Store) -> None:
    ws = wb.create_sheet("Excluded")
    ws.append(["Domain Key", "Company Name", "Exclusion Reason", "Excluded At", "Run ID"])
    for rec in store.conn.execute(
        "SELECT * FROM exclusions ORDER BY excluded_at DESC LIMIT 20000"
    ).fetchall():
        ws.append([
            rec["domain_key"], rec["company_name"], rec["reason"],
            rec["excluded_at"], rec["run_id"],
        ])
    _style_header(ws, [30, 34, 60, 22, 24])


def _sheet_run_report(wb: Workbook, report: dict[str, Any], store: Store) -> None:
    ws = wb.create_sheet("Run_Report")
    ws.append(["Metric", "Value"])

    scalar_first = [
        "run_id", "records_checked", "approved_count", "avg_completeness_pct",
        "duplicate_pairs", "duplicate_rate_pct", "links_checked", "links_broken",
        "approved_missing_level1_evidence", "approved_missing_rationale",
        "score_tier_mismatches", "passed",
    ]
    for key in scalar_first:
        if key in report:
            ws.append([key.replace("_", " ").title(), report[key]])

    ws.append([])
    ws.append(["Database Totals", ""])
    for key, value in store.counts().items():
        ws.append([key.replace("_", " ").title(), value])

    for section in ("tier_distribution", "segment_distribution", "state_distribution"):
        data = report.get(section)
        if not data:
            continue
        ws.append([])
        ws.append([section.replace("_", " ").title(), ""])
        for key, value in data.items():
            ws.append([str(key), value])

    issues = report.get("issues") or []
    ws.append([])
    ws.append(["Quality Issues", len(issues)])
    for issue in issues:
        ws.append(["", issue])

    examples = report.get("duplicate_examples") or []
    if examples:
        ws.append([])
        ws.append(["Possible Duplicate Pairs", ""])
        for state, a, b in examples:
            ws.append([state, f"{a}  <->  {b}"])

    _style_header(ws, [40, 80], autofilter=False)


def _sheet_learning(wb: Workbook, store: Store, run_id: str) -> None:
    ws = wb.create_sheet("Learning_Log")
    ws.append(["Kind", "Observation", "Detail", "Created At"])
    for rec in store.conn.execute(
        "SELECT * FROM learning_log WHERE run_id=? ORDER BY id", (run_id,)
    ).fetchall():
        ws.append([rec["kind"], rec["observation"], rec["detail"], rec["created_at"]])
    _style_header(ws, [20, 50, 70, 22])


def _cell(row: Any, key: str) -> Any:
    value = row[key]
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value)[:32000]
    if isinstance(value, str):
        # Excel hard-caps a cell at 32,767 characters.
        return value[:32000]
    return value
