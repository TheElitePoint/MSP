"""The master tracker — one accumulating workbook, exactly two sheets.

    Companies     one row per qualified account, with the date it was added
    Key_Contacts  the US decision-makers mapped inside those accounts

Rebuilt from SQLite on every run rather than appended to. The database is
already the accumulator — `first_seen_run`, `first_seen_at`, `last_updated_at`
and `times_seen` hold everything "keep adding data" needs, transactionally —
so the workbook is a view of it. Treating it as a view rather than as the
system of record means a crashed write cannot destroy the accumulated history,
ordering and tier colouring stay correct as scores move, and the whole file can
be regenerated at any time with `./run.py --export-only`.

The consequence to know about: edits made inside the spreadsheet are discarded
on the next run. Operator notes belong in the database (`review_notes`,
`next_action`), which do surface here as columns.

"Date Added" answers "what is new" without a second file: it is the immutable
first-seen date, and the `New?` column flags rows from this specific run.
"""

from __future__ import annotations

import logging
import os
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .export import HEADER_FILL, HEADER_FONT, TIER_FILLS, _cell, _style_header
from .store import Store

log = logging.getLogger("pipeline.master")

# (db column, header, width). Trimmed from the 59-column diagnostic tracker to
# the fields a person actually works from; everything dropped stays queryable
# in SQLite.
MASTER_COMPANY_COLUMNS: list[tuple[str, str, int]] = [
    ("_date_added", "Date Added", 12),
    ("_last_updated", "Last Updated", 12),
    ("_is_new", "New?", 7),
    # `Date Added` alone cannot separate two runs on the same day. The run id
    # carries date AND time (run_YYYYMMDD_HHMMSS), so filtering this column
    # isolates exactly one run's intake.
    ("first_seen_run", "Run ID", 22),
    ("tier", "Tier", 7),
    ("score_total", "Score", 8),
    ("status", "Status", 11),
    ("company_name", "Company Name", 34),
    ("website", "Website", 34),
    ("linkedin_url", "LinkedIn URL", 30),
    ("headquarters_city", "HQ City", 18),
    ("headquarters_state", "HQ State", 9),
    ("country", "Country", 14),
    ("phone", "Phone", 16),
    ("employee_range", "Employee Range", 14),
    ("entity_type", "Entity Type", 18),
    ("segment", "Segment", 28),
    ("services", "Services", 45),
    ("customer_industries", "Customer Industries", 28),
    ("company_description", "Company Description", 50),
    ("fit_reason", "Fit Reason", 70),
    ("operational_signals", "Operational Signals", 40),
    ("intent_score", "Intent Score (20)", 14),
    ("intent_summary", "Intent Summary", 50),
    ("timing_signal", "Timing Signal", 40),
    ("open_roles", "Open Ops Roles", 13),
    ("tools_found", "Tools Found", 40),
    ("score_fit", "Score: Fit (40)", 13),
    ("score_need", "Score: Need (25)", 14),
    ("score_timing", "Score: Timing (20)", 15),
    ("score_access", "Score: Access (15)", 15),
    ("_contacts", "Contacts Mapped", 15),
    ("next_action", "Next Action", 38),
    ("owner", "Owner", 12),
    ("review_notes", "Review Notes", 28),
    ("discovery_source", "Discovery Source", 28),
    ("discovery_query", "Discovery Query", 38),
    ("times_seen", "Times Seen", 11),
    ("id", "ID", 6),
]

MASTER_CONTACT_COLUMNS: list[tuple[str, str, int]] = [
    ("_date_added", "Date Added", 12),
    ("_last_seen", "Last Seen", 12),
    ("_is_new", "New?", 7),
    ("first_seen_run", "Run ID", 22),
    ("last_seen_run", "Last Seen Run", 22),
    ("company_id", "Company ID", 11),
    ("_company_name", "Company Name", 30),
    ("_company_tier", "Company Tier", 12),
    ("_company_state", "Company HQ State", 15),
    ("full_name", "Person Name", 24),
    ("job_title", "Job Title", 32),
    ("role_type", "Role Type", 14),
    ("seniority", "Seniority", 14),
    ("linkedin_url", "LinkedIn URL", 40),
    ("email", "Email", 28),
    ("email_status", "Email Status", 22),
    ("location", "Location", 26),
    ("us_state", "US State", 9),
    ("relevance", "Relevance", 55),
    ("about", "Profile Summary", 50),
    ("match_method", "Matched By", 14),
    ("date_checked", "Date Checked", 13),
    ("reject_reason", "Notes", 22),
]

ROLE_FILLS = {
    "economic": "C6EFCE",
    "technical": "DDEBF7",
    "operational": "FFF2CC",
    "champion": "E2EFDA",
    "risk": "FCE4D6",
}


def build_master_workbook(store: Store, cfg: Any, run_id: str) -> Path:
    """Write (or rewrite) the master tracker. Returns the path written."""
    out_dir = Path(cfg.get("export.master_dir", "data/exports"))
    if not out_dir.is_absolute():
        out_dir = Path(__file__).resolve().parent.parent / out_dir
    out_dir.mkdir(parents=True, exist_ok=True)

    label = cfg.get("run.label", "run")
    filename = str(cfg.get("export.master_filename", "master_{label}.xlsx"))
    path = out_dir / filename.format(label=label)

    statuses = cfg.get("export.include_statuses", ["approved", "review"])
    tiers = cfg.get("export.include_tiers", ["A"])
    verified_only = bool(cfg.get("export.verified_only", True))

    where = (f"status IN ({', '.join('?' for _ in statuses)})"
             f" AND (tier IN ({', '.join('?' for _ in tiers)}) OR tier IS NULL)")
    if verified_only:
        where += " AND clean_verified = 1"
    rows = store.companies_where(
        where + " ORDER BY CASE tier WHEN 'A' THEN 1 WHEN 'B' THEN 2 WHEN 'C' THEN 3"
        " ELSE 4 END, score_total DESC",
        (*statuses, *tiers),
    )

    contact_counts = store.contact_counts("accepted")
    company_index = {int(r["id"]): r for r in rows}

    wb = Workbook()
    wb.remove(wb.active)
    _sheet_master_companies(wb, rows, contact_counts, run_id, cfg)
    contacts = _collect_contacts(store, company_index)
    _sheet_master_contacts(wb, contacts, company_index, run_id, cfg)

    written = _save_atomic(wb, path)
    log.info(
        "Master tracker: %d companies, %d contacts -> %s",
        len(rows), len(contacts), written,
    )
    return written


# -- sheets ------------------------------------------------------------------


def _sheet_master_companies(
    wb: Workbook, rows: list[Any], contact_counts: dict[int, int],
    run_id: str, cfg: Any,
) -> None:
    ws = wb.create_sheet("Companies")
    ws.append([header for _, header, _ in MASTER_COMPANY_COLUMNS])

    for row in rows:
        derived = {
            "_date_added": _date(row["first_seen_at"]),
            "_last_updated": _date(row["last_updated_at"]),
            "_is_new": "NEW" if row["first_seen_run"] == run_id else "",
            "_contacts": contact_counts.get(int(row["id"]), 0),
        }
        ws.append([
            derived[key] if key in derived else _cell(row, key)
            for key, _, _ in MASTER_COMPANY_COLUMNS
        ])

    _colour(ws, MASTER_COMPANY_COLUMNS, "tier", TIER_FILLS)
    _bold_new(ws, MASTER_COMPANY_COLUMNS)
    _hyperlink(ws, MASTER_COMPANY_COLUMNS, "website")
    _hyperlink(ws, MASTER_COMPANY_COLUMNS, "linkedin_url")
    _style_header(
        ws, [w for _, _, w in MASTER_COMPANY_COLUMNS],
        freeze=bool(cfg.get("export.freeze_header", True)),
        autofilter=bool(cfg.get("export.autofilter", True)),
    )


def _sheet_master_contacts(
    wb: Workbook, contacts: list[Any], company_index: dict[int, Any],
    run_id: str, cfg: Any,
) -> None:
    ws = wb.create_sheet("Key_Contacts")
    ws.append([header for _, header, _ in MASTER_CONTACT_COLUMNS])

    for contact in contacts:
        company = company_index.get(int(contact["company_id"]))
        derived = {
            "_date_added": _date(contact["first_seen_at"]),
            "_last_seen": _date(contact["last_updated_at"]),
            "_is_new": "NEW" if contact["first_seen_run"] == run_id else "",
            "_company_name": company["company_name"] if company else "",
            "_company_tier": company["tier"] if company else "",
            "_company_state": company["headquarters_state"] if company else "",
        }
        ws.append([
            derived[key] if key in derived else _cell(contact, key)
            for key, _, _ in MASTER_CONTACT_COLUMNS
        ])

    _colour(ws, MASTER_CONTACT_COLUMNS, "role_type", ROLE_FILLS, as_fill=True)
    _bold_new(ws, MASTER_CONTACT_COLUMNS)
    _hyperlink(ws, MASTER_CONTACT_COLUMNS, "linkedin_url")
    _style_header(
        ws, [w for _, _, w in MASTER_CONTACT_COLUMNS],
        freeze=bool(cfg.get("export.freeze_header", True)),
        autofilter=bool(cfg.get("export.autofilter", True)),
    )


def _collect_contacts(store: Store, company_index: dict[int, Any]) -> list[Any]:
    """Accepted contacts for exported companies, in buying-priority order."""
    if not company_index:
        return []
    marks = ", ".join("?" for _ in company_index)
    return store.conn.execute(
        f"SELECT c.* FROM contacts c JOIN companies co ON co.id = c.company_id"
        f" WHERE c.status='accepted' AND c.company_id IN ({marks})"
        " ORDER BY CASE co.tier WHEN 'A' THEN 1 WHEN 'B' THEN 2 WHEN 'C' THEN 3 ELSE 4 END,"
        " co.score_total DESC, c.company_id, c.rank_score DESC",
        tuple(company_index),
    ).fetchall()


# -- formatting --------------------------------------------------------------


def _index_of(columns: list[tuple[str, str, int]], key: str) -> int | None:
    for idx, (col, _, _) in enumerate(columns, start=1):
        if col == key:
            return idx
    return None


def _colour(ws: Any, columns: list[tuple[str, str, int]], key: str,
            fills: dict[str, Any], as_fill: bool = False) -> None:
    col = _index_of(columns, key)
    if col is None:
        return
    from openpyxl.styles import PatternFill
    for excel_row in range(2, ws.max_row + 1):
        cell = ws.cell(row=excel_row, column=col)
        value = str(cell.value or "")
        fill = fills.get(value)
        if fill is not None:
            cell.fill = PatternFill("solid", fgColor=fill) if as_fill else fill
        cell.alignment = Alignment(horizontal="center")


def _bold_new(ws: Any, columns: list[tuple[str, str, int]]) -> None:
    col = _index_of(columns, "_is_new")
    if col is None:
        return
    for excel_row in range(2, ws.max_row + 1):
        cell = ws.cell(row=excel_row, column=col)
        if cell.value == "NEW":
            cell.font = Font(bold=True, color="1F7A1F")
        cell.alignment = Alignment(horizontal="center")


def _hyperlink(ws: Any, columns: list[tuple[str, str, int]], key: str) -> None:
    col = _index_of(columns, key)
    if col is None:
        return
    for excel_row in range(2, ws.max_row + 1):
        cell = ws.cell(row=excel_row, column=col)
        if cell.value and str(cell.value).startswith("http"):
            cell.hyperlink = str(cell.value)
            cell.font = Font(color="0563C1", underline="single")


def _date(value: Any) -> str:
    """YYYY-MM-DD as text, so Excel cannot reinterpret it by locale."""
    if not value:
        return ""
    return str(value)[:10]


def _save_atomic(wb: Workbook, path: Path) -> Path:
    """Write via a temp file so a crash cannot leave a half-written master.

    If the operator has the workbook open in Excel the replace fails; that is
    not worth aborting a whole run over, so it falls back to a stamped filename
    and says so loudly.
    """
    fd, tmp_name = tempfile.mkstemp(suffix=".xlsx", dir=str(path.parent))
    os.close(fd)
    tmp = Path(tmp_name)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
        return path
    except (PermissionError, OSError) as exc:
        stamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
        fallback = path.with_name(f"{path.stem}_{stamp}{path.suffix}")
        try:
            os.replace(tmp, fallback)
        except OSError:
            wb.save(fallback)
            tmp.unlink(missing_ok=True)
        log.warning(
            "Could not write %s (%s) — is it open in Excel? Wrote %s instead.",
            path.name, exc, fallback.name,
        )
        return fallback
